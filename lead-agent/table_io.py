"""Liest Excel-/CSV-Listen (z.B. eine Firmenliste, die Sebastian in den Vault
legt) als Zeilen-Dicts - Eingabe für check_companies/save_prospect-Läufe
("gleiche diese Liste mit Close ab"). Nur Dateien unterhalb des Vault-Mounts
bzw. des eigenen Export-Ordners, kein freier Dateisystemzugriff."""
import csv
from pathlib import Path

from config import get_settings
from export_leads import EXPORTS_DIR

MAX_ROWS = 1000
ALLOWED_SUFFIXES = {".xlsx", ".csv"}


def _allowed_roots() -> list[Path]:
    return [get_settings().vault_path.resolve(), Path(EXPORTS_DIR).resolve()]


def read_table(path: str, sheet: str = "", max_rows: int = 500) -> dict:
    raw = (path or "").strip()
    if not raw:
        return {"ok": False, "error": "Kein Pfad angegeben."}
    candidate = Path(raw)
    if not candidate.is_absolute():
        candidate = get_settings().vault_path / candidate
    try:
        resolved = candidate.resolve(strict=True)
    except (FileNotFoundError, OSError):
        return {"ok": False, "error": f"Datei nicht gefunden: {raw}"}
    if not any(root == resolved or root in resolved.parents for root in _allowed_roots()):
        return {"ok": False, "error": "Nur Dateien im Vault sind lesbar."}
    if resolved.suffix.lower() not in ALLOWED_SUFFIXES:
        return {"ok": False, "error": f"Nur {', '.join(sorted(ALLOWED_SUFFIXES))} unterstützt."}

    limit = max(1, min(int(max_rows or 500), MAX_ROWS))
    if resolved.suffix.lower() == ".csv":
        with resolved.open(encoding="utf-8-sig", newline="") as f:
            sample = f.read(4096)
            f.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            table = [list(r) for r in csv.reader(f, dialect)]
        sheets = []
        used_sheet = ""
    else:
        from openpyxl import load_workbook

        wb = load_workbook(resolved, read_only=True, data_only=True)
        sheets = wb.sheetnames
        used_sheet = sheet if sheet in sheets else sheets[0]
        table = [["" if c is None else c for c in row] for row in wb[used_sheet].iter_rows(values_only=True)]
        wb.close()

    table = [r for r in table if any(str(c).strip() for c in r)]
    if not table:
        return {"ok": False, "error": "Die Datei enthält keine Zeilen."}
    header = [str(h).strip() or f"Spalte{i + 1}" for i, h in enumerate(table[0])]
    rows = [{h: (str(v).strip() if v is not None else "") for h, v in zip(header, r)} for r in table[1:limit + 1]]
    return {
        "ok": True, "datei": resolved.name, "blatt": used_sheet, "blaetter": sheets,
        "spalten": header, "anzahl_zeilen_gesamt": len(table) - 1, "anzahl_zeilen": len(rows),
        "gekuerzt": len(table) - 1 > len(rows), "zeilen": rows,
    }
