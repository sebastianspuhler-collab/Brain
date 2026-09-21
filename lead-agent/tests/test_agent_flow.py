"""Regressions-/Integrationstests für die 2026-09-21-Umbauten: Statusfilter,
Spaltenexport, Tabellenexport, Excel-Einlesen, Domainlogik."""
import csv

import pytest

import combined_leads
import dedup
import export_leads
import prospects
import table_io
import vault_leads


def test_status_filter_matches_vault_status_of_leads_that_are_linked_to_close(vault, close):
    """Früher wurde status=neu als Close-Query geschickt (Close kennt 'neu' nicht)
    und warf damit ALLE Vault-Leads mit close_lead_id aus dem Ergebnis."""
    r = prospects.save_prospect("Verknuepft GmbH", website="verknuepft.de")
    rows = combined_leads.get_combined_leads({"status": "neu"})
    assert [x["close_lead_id"] for x in rows] == [r["close_lead_id"]]
    assert rows[0]["firma"] == "Verknuepft GmbH"       # Close-Anzeigename, nicht der Dateiname


def test_status_filter_also_matches_close_status(vault, close):
    close.add_lead("Termin AG", status="Termin vereinbart")
    close.add_lead("Kalt AG")
    rows = combined_leads.get_combined_leads({"status": "Termin vereinbart"})
    assert [x["firma"] for x in rows] == ["Termin AG"]
    assert rows[0]["close_status"] == "Termin vereinbart"


def test_region_and_branche_filters_use_close_fields(vault, close):
    close.add_lead("Kassel AG", custom={"Branche": "Elektrotechnik"}, addresses=[{"city": "Kassel"}])
    close.add_lead("Berlin AG", custom={"Branche": "Handel"}, addresses=[{"city": "Berlin"}])
    assert [x["firma"] for x in combined_leads.get_combined_leads({"region": "kassel"})] == ["Kassel AG"]
    assert [x["firma"] for x in combined_leads.get_combined_leads({"branche": "elektro"})] == ["Kassel AG"]


def test_close_outage_is_flagged_in_meta(vault, close):
    vault_leads.write_prospect("Nur Vault GmbH")
    close.down = True
    rows, meta = combined_leads.get_combined_leads_with_meta({})
    assert len(rows) == 1 and meta["close_verfuegbar"] is False and meta["close_fehler"]


def test_export_with_selected_columns_including_new_fields(vault, close, tmp_path, monkeypatch):
    monkeypatch.setattr(export_leads, "EXPORTS_DIR", tmp_path / "exports")
    prospects.save_prospect("Export GmbH", website="export-gmbh.de", ort="Ulm", branche="Elektro", quellen="https://export-gmbh.de", quelle="Recherche")
    res = export_leads.export_leads({}, "csv", "firma,website,ort,branche")
    assert res["ok"] and res["spalten"] == ["Firma", "Website", "Ort", "Branche"]
    rows = list(csv.reader((tmp_path / "exports" / res["filename"]).open(encoding="utf-8-sig")))
    assert rows[1] == ["Export GmbH", "https://export-gmbh.de", "Ulm", "Elektro"]


def test_export_rejects_unknown_column(vault, close, tmp_path, monkeypatch):
    monkeypatch.setattr(export_leads, "EXPORTS_DIR", tmp_path / "exports")
    res = export_leads.export_leads({}, "csv", "firma,gibt_es_nicht")
    assert res["ok"] is False and "gibt_es_nicht" in res["error"]


def test_export_table_writes_arbitrary_rows_and_flattens_lists(tmp_path, monkeypatch):
    import openpyxl

    monkeypatch.setattr(export_leads, "EXPORTS_DIR", tmp_path / "exports")
    res = export_leads.export_table(
        [{"firma": "A", "treffer": ["x", "y"], "note": None}, {"firma": "B", "treffer": [], "note": "ok"}],
        None, "xlsx", "Abgleich Test!",
    )
    assert res["ok"] and "Abgleich_Test" in res["filename"] and res["anzahl_zeilen"] == 2
    ws = openpyxl.load_workbook(tmp_path / "exports" / res["filename"]).active
    assert [c.value for c in ws[1]] == ["firma", "treffer", "note"]
    assert [c.value for c in ws[2]] == ["A", "x, y", None]   # leere Zelle
    assert export_leads.export_table([], None)["ok"] is False


def test_read_table_reads_xlsx_and_csv_inside_vault_only(vault, tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Firma", "Website"])
    ws.append(["Muster GmbH", "muster.de"])
    ws.append([None, None])
    ws.append(["Beispiel AG", None])
    (vault / "_inbox").mkdir()
    wb.save(vault / "_inbox" / "liste.xlsx")
    (vault / "_inbox" / "liste.csv").write_text("Firma;Website\nCSV GmbH;csv.de\n", encoding="utf-8")

    r = table_io.read_table("_inbox/liste.xlsx")
    assert r["ok"] and r["spalten"] == ["Firma", "Website"] and [z["Firma"] for z in r["zeilen"]] == ["Muster GmbH", "Beispiel AG"]
    assert table_io.read_table("_inbox/liste.csv")["zeilen"] == [{"Firma": "CSV GmbH", "Website": "csv.de"}]
    assert table_io.read_table("../../etc/passwd")["ok"] is False
    outside = tmp_path.parent / "fremd.csv"
    outside.write_text("a\n1\n")
    assert table_io.read_table(str(outside))["ok"] is False


@pytest.mark.parametrize("value,expected", [
    ("https://www.Beispiel-Technik.de/kontakt", "beispiel-technik.de"),
    ("info@shop.muster.de", "shop.muster.de"),
    ("hans@gmail.com", ""),
    ("kein host", ""),
])
def test_domain_of(value, expected):
    assert dedup.domain_of(value) == expected


def test_domain_match_ignores_freemail_and_handles_subdomains():
    assert dedup.domains_match("shop.muster.de", "muster.de")
    assert not dedup.domains_match("muster.de", "muster-neu.de")
